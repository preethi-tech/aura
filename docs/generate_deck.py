"""Generate the Aura hackathon presentation (PPTX) with diagrams.

Run:  python docs/generate_deck.py
Output: docs/Aura_Presentation.pptx

Pure python-pptx; no external template needed. Diagrams are drawn with native
shapes/connectors so the deck is fully editable in PowerPoint / Google Slides.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# --- Palette (matches the app) ---------------------------------------------
INK = RGBColor(0x1A, 0x1D, 0x23)
BG = RGBColor(0x0D, 0x11, 0x17)
SURFACE = RGBColor(0x16, 0x1B, 0x22)
CARD = RGBColor(0x1C, 0x23, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0x8B, 0x94, 0x9E)
ACCENT = RGBColor(0x4F, 0x6E, 0xF7)
ACCENT2 = RGBColor(0x6D, 0x8A, 0xFF)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
YELLOW = RGBColor(0xEA, 0xB3, 0x08)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
RED = RGBColor(0xEF, 0x44, 0x44)
LIGHT = RGBColor(0xF5, 0xF6, 0xF8)

# 16:9
EMU_W = 12192000
EMU_H = 6858000

prs = Presentation()
prs.slide_width = Emu(EMU_W)
prs.slide_height = Emu(EMU_H)
BLANK = prs.slide_layouts[6]


def inch(v):
    return Emu(int(v * 914400))


def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color=BG):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space=1.06):
    tb = s.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, 18, WHITE, False)]
    first = True
    for item in runs:
        text, size, color, bold = (item + (False,))[:4] if len(item) < 4 else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = space
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Segoe UI"
    return tb


def shape_text(sp, text, size=14, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = "Segoe UI"
    return sp


def connector(s, x1, y1, x2, y2, color=ACCENT, w=2.0, arrow=True):
    cn = s.shapes.add_connector(2, inch(x1), inch(y1), inch(x2), inch(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(w)
    if arrow:
        line_elem = cn.line._get_or_add_ln()
        from pptx.oxml.ns import qn
        tail = line_elem.makeelement(qn("a:tailEnd"),
                                     {"type": "triangle", "w": "med", "h": "med"})
        line_elem.append(tail)
    return cn


def kicker(s, text, color=ACCENT2):
    txt(s, 0.6, 0.45, 8, 0.4, [(text.upper(), 13, color, True)])


def title(s, text, y=0.8, size=34):
    txt(s, 0.6, y, 11, 1.0, [(text, size, WHITE, True)])


def footer(s, n):
    txt(s, 11.0, 6.35, 1.0, 0.3, [(f"{n}", 11, MUTE, False)], align=PP_ALIGN.RIGHT)
    txt(s, 0.6, 6.35, 4, 0.3, [("Aura \u2014 wellbeing early-warning", 11, MUTE, False)])


# ===========================================================================
# 1. TITLE
# ===========================================================================
s = slide(); bg(s)
box(s, 0, 0, 13.333, 7.5, fill=BG, shape=MSO_SHAPE.RECTANGLE)
# accent orb
o = box(s, 10.4, -1.2, 4, 4, fill=ACCENT, shape=MSO_SHAPE.OVAL)
o.fill.fore_color.rgb = ACCENT
o2 = box(s, 11.1, 4.6, 3.2, 3.2, fill=CARD, shape=MSO_SHAPE.OVAL)
txt(s, 0.9, 2.2, 9, 1.0, [("\u25c2  AURA", 20, ACCENT2, True)])
txt(s, 0.9, 2.75, 10.5, 1.6, [("A private early-warning companion", 46, WHITE, True),
                              ("for mental wellbeing", 46, WHITE, True)])
txt(s, 0.9, 4.75, 10, 0.8, [(
    "Notice sustained changes in your own patterns \u2014 and connect to "
    "support earlier.", 18, MUTE, False)])
txt(s, 0.9, 5.7, 11, 0.5, [(
    "Passive signals  \u2022  Explainable AI  \u2022  Micro-interventions  "
    "\u2022  Circle of Care  \u2022  Built on Google Cloud", 14, ACCENT2, True)])

# ===========================================================================
# 2. PROBLEM
# ===========================================================================
s = slide(); bg(s)
kicker(s, "The problem")
title(s, "People notice too late \u2014 and no one else notices at all")
pts = [
    ("Silent onset", "Depression and anxiety build gradually. By the time it's "
     "obvious, the person is already deep in it."),
    ("Manual tools fail", "Daily mood diaries demand effort exactly when energy "
     "and motivation are lowest \u2014 so people stop."),
    ("Diagnostic, not preventive", "Existing apps either diagnose (stigmatising) "
     "or react in crisis. Nothing supports quiet, early self-awareness."),
    ("The 'who notices' gap", "The people who could help often have no signal "
     "that something has shifted."),
]
x = 0.6
for i, (h, b) in enumerate(pts):
    cx = 0.6 + (i % 2) * 6.05
    cy = 2.0 + (i // 2) * 2.05
    box(s, cx, cy, 5.75, 1.85, fill=CARD, line=RED if i in (0, 3) else ORANGE, line_w=1.5)
    txt(s, cx + 0.3, cy + 0.22, 5.2, 0.5, [(h, 18, WHITE, True)])
    txt(s, cx + 0.3, cy + 0.75, 5.2, 1.0, [(b, 13.5, MUTE, False)])
footer(s, 2)

# ===========================================================================
# 3. SOLUTION
# ===========================================================================
s = slide(); bg(s)
kicker(s, "The solution")
title(s, "Aura turns weak everyday signals into an early, gentle heads-up")
txt(s, 0.6, 1.75, 11.5, 0.9, [(
    "A private, opt-in companion that fuses active check-ins and passive "
    "phone/wearable signals into one explainable 'Aura Index', personalised to "
    "your own baseline \u2014 then closes the loop with small actions and, only "
    "if you choose, your trusted circle.", 16, LIGHT, False)])
pillars = [
    ("Personalised", "Every signal is scored against YOUR history, never a "
     "population norm.", ACCENT),
    ("Sustained, not spiky", "One rough day never alarms; multi-signal drift "
     "over time does.", GREEN),
    ("Explainable", "Every score decomposes into named, human-readable "
     "contributions.", YELLOW),
    ("Private by design", "Runs offline-capable; the user owns and can export "
     "or delete everything.", ACCENT2),
]
for i, (h, b, col) in enumerate(pillars):
    cx = 0.6 + i * 3.0
    box(s, cx, 3.1, 2.8, 2.7, fill=CARD, line=col, line_w=1.5)
    box(s, cx + 0.3, 3.35, 0.5, 0.5, fill=col, shape=MSO_SHAPE.OVAL)
    txt(s, cx + 0.25, 4.0, 2.4, 0.6, [(h, 16.5, WHITE, True)])
    txt(s, cx + 0.25, 4.6, 2.4, 1.1, [(b, 12.5, MUTE, False)])
footer(s, 3)

# ===========================================================================
# 4. HOW IT WORKS (pipeline diagram)
# ===========================================================================
s = slide(); bg(s)
kicker(s, "How it works")
title(s, "From signals to a sustained, explainable index")
stages = [
    ("Inputs", "Journal + sleep,\nsocial, energy\n+ steps, screen time", ACCENT),
    ("Features", "Sentiment, self-focus,\nabsolutist &\nfuture-focus language", ACCENT2),
    ("Personal baseline", "Robust median / MAD\nover YOUR recent\nhistory (z-scores)", YELLOW),
    ("Fuse + smooth", "Weight & combine;\nrolling window enforces\n'sustained' drift", ORANGE),
    ("Aura Index 0\u2013100", "Tiered heads-up with\nper-signal 'why'", GREEN),
]
n = len(stages)
w = 2.15
gap = (12.13 - n * w) / (n - 1)
y = 2.5
for i, (h, b, col) in enumerate(stages):
    cx = 0.6 + i * (w + gap)
    box(s, cx, y, w, 1.9, fill=CARD, line=col, line_w=1.5)
    txt(s, cx + 0.12, y + 0.15, w - 0.24, 0.6, [(h, 14.5, col, True)], align=PP_ALIGN.CENTER)
    txt(s, cx + 0.12, y + 0.72, w - 0.24, 1.1, [(b, 11.5, MUTE, False)], align=PP_ALIGN.CENTER)
    if i < n - 1:
        connector(s, cx + w, y + 0.95, cx + w + gap, y + 0.95, color=WHITE, w=2.0)
box(s, 0.6, 4.9, 11.9, 1.2, fill=SURFACE, line=ACCENT, line_w=1.2)
txt(s, 0.85, 5.05, 11.4, 1.0, [
    ("Why this design: ", 14, ACCENT2, True),
    ("A single bad day produces a tiny, transient blip. Only when several weak "
     "signals drift together and stay drifted does the index climb into higher "
     "tiers \u2014 mirroring how real decline actually looks, while minimising "
     "false alarms.", 13.5, LIGHT, False),
], space=1.05)
footer(s, 4)

# ===========================================================================
# 5. ARCHITECTURE (Google Cloud)
# ===========================================================================
s = slide(); bg(s)
kicker(s, "Architecture")
title(s, "Cloud-native, built on Google Cloud")
# client
box(s, 0.6, 2.2, 2.6, 1.4, fill=CARD, line=ACCENT2, line_w=1.5)
txt(s, 0.7, 2.4, 2.4, 1.0, [("Browser UI", 15, WHITE, True),
                            ("Light/dark, charts,\nFirebase Auth", 11.5, MUTE, False)],
    align=PP_ALIGN.CENTER)
# cloud run
box(s, 4.3, 2.2, 3.0, 1.4, fill=SURFACE, line=ACCENT, line_w=2.0)
txt(s, 4.4, 2.35, 2.8, 1.1, [("Cloud Run", 16, ACCENT2, True),
                             ("FastAPI + static UI\n(Docker container)", 12, MUTE, False)],
    align=PP_ALIGN.CENTER)
connector(s, 3.2, 2.9, 4.3, 2.9, color=WHITE)
# services on the right
svcs = [
    ("Gemini API", "Sentiment + warm\nweekly summaries", ACCENT),
    ("Firestore", "Per-user cloud\npersistence", GREEN),
    ("Firebase Auth", "Google Sign-in,\nuser isolation", YELLOW),
    ("Cloud Logging", "Structured\nobservability", ORANGE),
]
for i, (h, b, col) in enumerate(svcs):
    cy = 0.9 + i * 1.35
    box(s, 8.5, cy, 3.4, 1.15, fill=CARD, line=col, line_w=1.3)
    txt(s, 8.65, cy + 0.12, 3.1, 0.5, [(h, 14, col, True)])
    txt(s, 8.65, cy + 0.55, 3.1, 0.55, [(b, 11, MUTE, False)])
    connector(s, 7.3, 2.9, 8.5, cy + 0.55, color=col, w=1.5)
# supporting
box(s, 4.3, 4.1, 3.0, 1.0, fill=CARD, line=MUTE, line_w=1.2)
txt(s, 4.4, 4.25, 2.8, 0.8, [("Secret Manager", 13.5, WHITE, True),
                             ("API keys \u2014 never in code", 11, MUTE, False)],
    align=PP_ALIGN.CENTER)
connector(s, 5.8, 3.6, 5.8, 4.1, color=MUTE, arrow=False)
box(s, 0.6, 4.1, 2.6, 1.0, fill=CARD, line=MUTE, line_w=1.2)
txt(s, 0.7, 4.25, 2.4, 0.8, [("SQLite fallback", 13.5, WHITE, True),
                             ("Fully offline mode", 11, MUTE, False)],
    align=PP_ALIGN.CENTER)
connector(s, 1.9, 3.6, 1.9, 4.1, color=MUTE, arrow=False)
txt(s, 0.6, 5.5, 11.5, 0.8, [(
    "One image, two modes: with keys it uses Gemini + Firestore + Cloud "
    "Logging; with none it runs 100% offline on SQLite with a deterministic "
    "analyzer. Same code, graceful degradation.", 13.5, LIGHT, False)])
footer(s, 5)

# ===========================================================================
# 6-9 FEATURES
# ===========================================================================
def feature_slide(n, kick, ttl, blurb, bullets, tag, tagcol):
    s = slide(); bg(s)
    kicker(s, kick)
    title(s, ttl)
    # tag chip
    chip = box(s, 0.6, 1.65, 2.4, 0.5, fill=tagcol)
    shape_text(chip, tag, 12.5, WHITE, True)
    txt(s, 0.6, 2.35, 6.4, 1.4, [(blurb, 15.5, LIGHT, False)])
    for i, (h, b) in enumerate(bullets):
        cy = 3.7 + i * 0.95
        box(s, 0.6, cy, 6.4, 0.85, fill=CARD, line=tagcol, line_w=1.2)
        txt(s, 0.8, cy + 0.1, 6.0, 0.7, [(h + "  ", 13.5, tagcol, True),
                                         (b, 12.5, MUTE, False)], space=1.0)
    # right visual panel (mock)
    box(s, 7.4, 1.9, 5.3, 4.2, fill=SURFACE, line=tagcol, line_w=1.5)
    return s


s = feature_slide(
    6, "Feature 1", "Passive Signal Integration",
    "Reduce manual effort and increase reliability by pulling sleep + activity "
    "from Google Fit and deriving passive phone signals \u2014 without ever "
    "reading content.",
    [("Google Fit", "steps, active minutes & sleep via the Fitness REST API"),
     ("Phone signal", "daily screen-time as a digital-phenotyping marker"),
     ("Augments, never fabricates", "passive data only enriches real check-in days"),
     ("Graceful demo", "realistic simulation when no OAuth token is present")],
    "wearable + phone", ACCENT)
txt(s, 7.7, 2.15, 4.7, 0.5, [("One-tap: 'Sync Google Fit'", 15, WHITE, True)])
for i, (lbl, val, col) in enumerate([
        ("Steps", "8,500 \u2192 2,200", GREEN),
        ("Active min", "45 \u2192 8", YELLOW),
        ("Screen time", "200 \u2192 430 min", RED)]):
    cy = 2.8 + i * 1.05
    box(s, 7.7, cy, 4.7, 0.9, fill=CARD, line=col, line_w=1.2)
    txt(s, 7.9, cy + 0.16, 2.5, 0.6, [(lbl, 14, WHITE, True)])
    txt(s, 10.0, cy + 0.16, 2.3, 0.6, [(val, 14, col, True)], align=PP_ALIGN.RIGHT)
footer(s, 6)

s = feature_slide(
    7, "Feature 2", "Explainable AI Insights",
    "Beyond a number: Aura explains itself. Gemini turns grounded facts into a "
    "warm weekly narrative and surfaces recurring cycles \u2014 reasoning, not "
    "just scoring.",
    [("Weekly summary", "'sleep 2h below baseline for 5 days, with less contact'"),
     ("Recurring cycles", "weekday-vs-weekend and your toughest day, detected"),
     ("Numbers stay real", "Gemini rephrases; it never invents the figures"),
     ("Works offline", "deterministic narrative when Gemini is off")],
    "Gemini-powered", ACCENT2)
txt(s, 7.7, 2.15, 4.7, 0.5, [("Weekly Summary", 15, WHITE, True)])
box(s, 7.7, 2.7, 4.7, 1.6, fill=CARD, line=ACCENT2, line_w=1.2)
txt(s, 7.9, 2.82, 4.3, 1.5, [(
    "\u201cYour index rose this week (avg 72/100). Sleep averaged 5.1h, about "
    "2h below your usual, alongside less social contact. Be gentle with "
    "yourself \u2014 small steps count.\u201d", 12.5, LIGHT, False)])
txt(s, 7.7, 4.45, 4.7, 0.4, [("Recurring pattern", 15, WHITE, True)])
box(s, 7.7, 4.9, 4.7, 1.0, fill=CARD, line=YELLOW, line_w=1.2)
txt(s, 7.9, 5.05, 4.3, 0.8, [(
    "Your index tends to run higher on weekends; Saturday is often your "
    "toughest day.", 12.5, MUTE, False)])
footer(s, 7)

s = feature_slide(
    8, "Feature 3", "Micro-Intervention Suggestions",
    "When the index rises, Aura offers ONE small, evidence-based action mapped "
    "to the signal driving it \u2014 closing the loop from detection to action, "
    "never overstepping into medical advice.",
    [("Targeted", "action matches the top contributing signal"),
     ("Tiny & doable", "2\u201310 minute steps (box breathing, a short walk)"),
     ("Evidence-based", "behavioural activation & sleep-hygiene grounding"),
     ("Kind at baseline", "a 'keep it up' note when things are stable")],
    "detection \u2192 action", GREEN)
txt(s, 7.7, 2.15, 4.7, 0.5, [("Try this now", 15, WHITE, True)])
for i, (t, d, a) in enumerate([
        ("Wind-down cue", "tonight", "Screens away 30 min before a fixed lights-out."),
        ("Box breathing", "2 min", "In 4, hold 4, out 4, hold 4 \u2014 repeat."),
        ("One small reconnect", "2 min", "Send a one-line hello to someone you trust.")]):
    cy = 2.7 + i * 1.1
    box(s, 7.7, cy, 4.7, 0.95, fill=CARD, line=GREEN, line_w=1.2)
    txt(s, 7.9, cy + 0.1, 3.2, 0.4, [(t, 13.5, GREEN, True)])
    chip = box(s, 11.15, cy + 0.12, 1.05, 0.35, fill=GREEN)
    shape_text(chip, d, 10.5, WHITE, True)
    txt(s, 7.9, cy + 0.5, 4.3, 0.4, [(a, 11.5, MUTE, False)])
footer(s, 8)

s = feature_slide(
    9, "Feature 4", "Circle of Care", 
    "Solve the 'who notices' gap with privacy by default. Pre-select 1\u20132 "
    "trusted people; at the Urgent tier Aura privately suggests you reach out "
    "with a ready, generic message.",
    [("Opt-in only", "no contacts exist unless you add them"),
     ("Never auto-sends", "Aura prepares a message; you choose to send it"),
     ("Data minimisation", "generic nudge \u2014 no scores, signals or journal"),
     ("You're in control", "human-in-the-loop at every step")],
    "opt-in + private", RED)
txt(s, 7.7, 2.15, 4.7, 0.5, [("Suggested nudge (you send)", 14.5, WHITE, True)])
box(s, 7.7, 2.75, 4.7, 1.5, fill=CARD, line=RED, line_w=1.2)
txt(s, 7.9, 2.9, 4.3, 1.3, [(
    "\u201cHi Sam, quick hello. I've been using a wellbeing app and set you as "
    "someone I trust. I'd appreciate a check-in when you have a moment.\u201d",
    12.5, LIGHT, False)])
box(s, 7.7, 4.45, 4.7, 1.45, fill=SURFACE, line=ACCENT2, line_w=1.2)
txt(s, 7.9, 4.58, 4.3, 1.3, [
    ("Privacy guarantee  ", 12.5, ACCENT2, True),
    ("Aura never contacts anyone for you and never shares your data. The nudge "
     "is a private suggestion \u2014 nothing more.", 12, MUTE, False)])
footer(s, 9)

# ===========================================================================
# 10. TECH STACK & SERVICES
# ===========================================================================
s = slide(); bg(s)
kicker(s, "Under the hood")
title(s, "Tech stack & Google services")
cols = [
    ("Backend", ACCENT, ["Python 3.12", "FastAPI + Uvicorn", "Deterministic analyzer",
                         "Pydantic schemas", "65 passing tests"]),
    ("Frontend", ACCENT2, ["Vanilla JS (no framework)", "Chart.js visualisations",
                          "Light / dark themes", "Responsive UI", "Firebase Auth SDK"]),
    ("Google Cloud", GREEN, ["Cloud Run (serverless)", "Gemini API", "Firebase Firestore",
                            "Firebase Auth", "Cloud Logging + Secret Mgr"]),
]
for i, (h, col, items) in enumerate(cols):
    cx = 0.6 + i * 4.05
    box(s, cx, 1.9, 3.8, 4.1, fill=CARD, line=col, line_w=1.5)
    bar = box(s, cx, 1.9, 3.8, 0.65, fill=col, shape=MSO_SHAPE.RECTANGLE)
    shape_text(bar, h, 17, WHITE, True)
    for j, it in enumerate(items):
        txt(s, cx + 0.35, 2.8 + j * 0.6, 3.2, 0.5,
            [("\u2022  ", 14, col, True), (it, 13.5, LIGHT, False)], space=1.0)
footer(s, 10)

# ===========================================================================
# 11. PRIVACY & ETHICS
# ===========================================================================
s = slide(); bg(s)
kicker(s, "Responsible by design", GREEN)
title(s, "Privacy, safety & ethics are first-class")
items = [
    ("Not a medical device", "No diagnosis, treatment or prediction claims \u2014 "
     "a self-awareness aid, explicitly.", RED),
    ("Safety never depends on ML", "Explicit crisis language always surfaces "
     "helplines immediately, regardless of the score.", ORANGE),
    ("Data minimisation", "Passive signals are aggregate counts; we never read "
     "message content. Contacts get generic nudges only.", ACCENT),
    ("User ownership", "One-tap export or delete-everything. Offline SQLite mode "
     "means data can stay fully on-device.", GREEN),
    ("Human in the loop", "Aura suggests; the person always decides and acts.", ACCENT2),
    ("Explainable, not a black box", "Every score decomposes into named, "
     "inspectable contributions.", YELLOW),
]
for i, (h, b, col) in enumerate(items):
    cx = 0.6 + (i % 2) * 6.05
    cy = 1.95 + (i // 2) * 1.42
    box(s, cx, cy, 5.75, 1.28, fill=CARD, line=col, line_w=1.3)
    txt(s, cx + 0.28, cy + 0.14, 5.2, 0.45, [(h, 15, WHITE, True)])
    txt(s, cx + 0.28, cy + 0.6, 5.2, 0.65, [(b, 12, MUTE, False)])
footer(s, 11)

# ===========================================================================
# 12. VALIDATION / RESULTS
# ===========================================================================
s = slide(); bg(s)
kicker(s, "Does it actually work?")
title(s, "Validated against PHQ-9 / GAD-7")
txt(s, 0.6, 1.75, 11.5, 0.8, [(
    "A built-in evaluation harness compares the passive Aura Index against "
    "periodic PHQ-9 / GAD-7 self-reports on the demo timeline \u2014 measuring "
    "correlation and, crucially, how much EARLIER it flags a shift.", 15, LIGHT, False)])
metrics = [
    ("Correlation", "Aura Index tracks\nPHQ-9 & GAD-7", ACCENT),
    ("Lead time", "Flags a shift days\nBEFORE the questionnaire", GREEN),
    ("Precision / Recall / F1", "Day-level detection\nquality reported", YELLOW),
    ("65 tests", "Backend fully\nunit + API tested", ACCENT2),
]
for i, (h, b, col) in enumerate(metrics):
    cx = 0.6 + i * 3.0
    box(s, cx, 2.9, 2.8, 2.2, fill=CARD, line=col, line_w=1.5)
    txt(s, cx + 0.2, 3.15, 2.4, 0.8, [(h, 15.5, col, True)], align=PP_ALIGN.CENTER)
    txt(s, cx + 0.2, 4.0, 2.4, 1.0, [(b, 12.5, MUTE, False)], align=PP_ALIGN.CENTER)
txt(s, 0.6, 5.5, 11.5, 0.7, [(
    "Metrics are illustrative on synthetic/self-reported data; real clinical "
    "validation would require an IRB-approved prospective study.", 12, MUTE, False)])
footer(s, 12)

# ===========================================================================
# 13. ROADMAP / CLOSE
# ===========================================================================
s = slide(); bg(s)
box(s, 0, 0, 13.333, 7.5, fill=BG, shape=MSO_SHAPE.RECTANGLE)
o = box(s, -1.4, 4.2, 4, 4, fill=CARD, shape=MSO_SHAPE.OVAL)
txt(s, 0.9, 0.9, 11, 1.0, [("What's next", 30, WHITE, True)])
road = [
    ("Live Google Fit OAuth", "swap the simulator for a full consented OAuth2 flow"),
    ("Vertex AI patterns", "richer seasonal / cyclical detection at scale"),
    ("Pub/Sub + Functions", "event-driven nudges and reminders"),
    ("BigQuery (opt-in, anonymised)", "cohort insights for research partners"),
    ("Native mobile", "background passive collection with on-device privacy"),
]
for i, (h, b) in enumerate(road):
    cy = 1.9 + i * 0.72
    box(s, 0.9, cy, 0.16, 0.5, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
    txt(s, 1.25, cy, 10.5, 0.6, [(h + "  \u2014  ", 15, ACCENT2, True),
                                 (b, 14, LIGHT, False)], space=1.0)
box(s, 0.9, 5.75, 11.5, 1.1, fill=SURFACE, line=ACCENT, line_w=1.5)
txt(s, 1.15, 5.92, 11.0, 0.9, [
    ("Aura  ", 16, ACCENT2, True),
    ("\u2014 weak signals, noticed early, explained kindly, and acted on "
     "together. Private by default, powered by Google Cloud.", 15, WHITE, False)])
footer(s, 13)

out = Path(__file__).resolve().parent / "Aura_Presentation.pptx"
prs.save(str(out))
print(f"Saved {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
